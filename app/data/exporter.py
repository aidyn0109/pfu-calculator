"""Экспорт результатов расчёта в Excel (openpyxl).

Числа выгружаются числовым форматом с 2 знаками (не текстом), как требует
раздел 7.3 CLAUDE.md.
"""

from __future__ import annotations

import json
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.config import BRAND_COLOR_PRIMARY, DISPLAY_DECIMALS, KIND_MDE
from app.db.calculations import kind_of, params_of, years_of
from app.db.models import Calculation

# Подписи параметров MDE в шапке выгрузки (порядок — как в форме на странице).
_PARAM_LABELS: list[tuple[str, str]] = [
    ("revenue_cap", "Кап дохода"),
    ("taxes_threshold", "Налоги: порог"),
    ("taxes_step", "Налоги: шаг"),
    ("taxes_factor", "Налоги: коэффициент"),
    ("taxes_cap", "Налоги: кап"),
    ("payroll_threshold", "ФОТ: порог"),
    ("payroll_step", "ФОТ: шаг"),
    ("payroll_factor", "ФОТ: коэффициент"),
    ("payroll_cap", "ФОТ: кап"),
]

_NUM_FMT = "0." + "0" * DISPLAY_DECIMALS

# Заголовки и ключи результата (порядок столбцов экспорта).
_COLUMNS: list[tuple[str, str, bool]] = [
    ("БИН", "bin", False),
    ("Наименование", "name", False),
    ("Годы расчёта", "years_used", False),
    ("Доход, %", "revenue_percent", True),
    ("Доход (индикатор)", "revenue_indicator", True),
    ("Ограничение дохода", "revenue_cap_applied", False),
    ("Налоги, %", "taxes_percent", True),
    ("Налоги (индикатор)", "taxes_indicator", True),
    ("Ограничение налогов", "taxes_cap_applied", False),
    ("ФОТ, %", "payroll_percent", True),
    ("ФОТ (индикатор)", "payroll_indicator", True),
    ("Ограничение ФОТ", "payroll_cap_applied", False),
    ("ПФУ (до ограничения)", "pfu_raw", True),
    ("ПФУ итог", "pfu_final", True),
    ("Ограничение ПФУ", "pfu_cap_applied", False),
    ("Предупреждения / ошибка", "warnings", False),
]


def _cell_value(result: dict[str, Any], key: str) -> Any:
    """Привести значение результата к виду для ячейки Excel."""
    if key == "years_used":
        return ", ".join(str(y) for y in result.get("years_used", []))
    if key == "warnings":
        parts = list(result.get("warnings", []))
        if result.get("error"):
            parts.append(result["error"])
        return "; ".join(parts)
    value = result.get(key)
    if isinstance(value, bool):
        return "да" if value else "нет"
    return value


def build_export(calc: Calculation) -> bytes:
    """Сформировать .xlsx с результатами расчёта. Возвращает байты файла."""
    results: list[dict[str, Any]] = json.loads(calc.results)

    wb = Workbook()
    ws = wb.active
    ws.title = "ПФУ"

    header_fill = PatternFill("solid", fgColor=BRAND_COLOR_PRIMARY.lstrip("#"))
    header_font = Font(bold=True, color="FFFFFF")

    # Шапка с метаданными расчёта.
    is_mde = kind_of(calc) == KIND_MDE
    title = "Расчёт MDE" if is_mde else "Расчёт"
    ws.append([f"{title} #{calc.id} от {calc.created_at}"])
    ws.append([f"Сумма: {calc.amount}", f"Сумма (МРП): {calc.amount_mrp}"])
    ws.append(["Годы расчёта: " + ", ".join(str(y) for y in years_of(calc))])

    # Для MDE — параметры формул, которые задал пользователь.
    params = params_of(calc)
    if params:
        ws.append(["Параметры расчёта:"])
        for key, label in _PARAM_LABELS:
            value = params.get(key)
            if key == "revenue_cap" and value is None:
                value = "авто по диапазону суммы"
            ws.append([label, value])

    ws.append([])

    header_row_idx = ws.max_row + 1
    ws.append([title for title, _, _ in _COLUMNS])
    for cell in ws[header_row_idx]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for result in results:
        row_values = [_cell_value(result, key) for _, key, _ in _COLUMNS]
        ws.append(row_values)
        row_idx = ws.max_row
        for col_idx, (_, _, is_number) in enumerate(_COLUMNS, start=1):
            cell = ws.cell(row=row_idx, column=col_idx)
            if is_number and isinstance(cell.value, (int, float)):
                cell.number_format = _NUM_FMT

    # Простейшая авто-ширина столбцов.
    for col_idx, (title, _, _) in enumerate(_COLUMNS, start=1):
        letter = ws.cell(row=header_row_idx, column=col_idx).column_letter
        ws.column_dimensions[letter].width = max(14, len(title) + 2)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
